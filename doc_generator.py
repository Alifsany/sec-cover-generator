import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

# Register TTF fonts for ReportLab if available
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

def find_asset(filename):
    p1 = os.path.join(BASE_DIR, filename)
    if os.path.exists(p1):
        return p1
    p2 = os.path.join(BASE_DIR, 'static', filename)
    if os.path.exists(p2):
        return p2
    return p1

ANTONIO_FONT = find_asset('Antonio-Bold.ttf')
COMFORTAA_FONT = find_asset('Comfortaa-Regular.ttf')
WATERMARK_IMG = find_asset('watermark_transparent.png')

try:
    if os.path.exists(ANTONIO_FONT):
        pdfmetrics.registerFont(TTFont('Antonio-Bold', ANTONIO_FONT))
    if os.path.exists(COMFORTAA_FONT):
        pdfmetrics.registerFont(TTFont('Comfortaa-Regular', COMFORTAA_FONT))
except Exception as e:
    print('Font registration error:', e)

def generate_docx(title, course, sname, sid, semester, tname, tdesignation, output_path):
    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.different_first_page_header_footer = False

    # Watermark in Header
    if os.path.exists(WATERMARK_IMG):
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hrun = hp.add_run()
        pic = hrun.add_picture(WATERMARK_IMG, width=Inches(5.55))
        
        inline = pic._inline
        cx = inline.extent.cx
        cy = inline.extent.cy
        doc_id = inline.docPr.id
        doc_name = inline.docPr.name
        top_emu = int(2.42 * 914400)

        anchor_xml = f"""
        <wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
                   distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240"
                   behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
          <wp:simplePos x="0" y="0"/>
          <wp:positionH relativeFrom="page"><wp:align>center</wp:align></wp:positionH>
          <wp:positionV relativeFrom="page"><wp:posOffset>{top_emu}</wp:posOffset></wp:positionV>
          <wp:extent cx="{cx}" cy="{cy}"/>
          <wp:effectExtent l="0" t="0" r="0" b="0"/>
          <wp:wrapNone/>
          <wp:docPr id="{doc_id}" name="{doc_name}"/>
          <wp:cNvGraphicFramePr/>
        </wp:anchor>
        """
        anchor = parse_xml(anchor_xml)
        graphic = inline.graphic
        anchor.append(graphic)
        inline.getparent().replace(inline, anchor)

    # Paragraph 1: SHYAMOLI ENGINEERING COLLEGE
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(14)
    p1.paragraph_format.space_after = Pt(6)
    p1.paragraph_format.line_spacing = 1.0
    r1 = p1.add_run("SHYAMOLI ENGINEERING COLLEGE")
    r1.font.name = "Antonio"
    r1.font.size = Pt(43)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0, 0, 0)
    r1._r.get_or_add_rPr().append(parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Antonio" w:hAnsi="Antonio" w:cs="Antonio"/>'))

    # Paragraph 2: A CONSTITUENT COLLEGE OF UNIVERSITY OF DHAKA
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(66)
    p2.paragraph_format.line_spacing = 1.0
    r2 = p2.add_run("A CONSTITUENT COLLEGE OF UNIVERSITY OF DHAKA")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(7, 122, 205)

    # Paragraph 3: DEPARTMENT OF
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(4)
    p3.paragraph_format.line_spacing = 1.15
    r3 = p3.add_run("DEPARTMENT OF")
    r3.font.name = "Comfortaa"
    r3.font.size = Pt(20)
    r3.font.bold = True
    r3._r.get_or_add_rPr().append(parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Comfortaa" w:hAnsi="Comfortaa" w:cs="Comfortaa"/>'))

    # Paragraph 4: COMPUTER SCIENCE AND ENGINEERING
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(0)
    p4.paragraph_format.space_after = Pt(48)
    p4.paragraph_format.line_spacing = 1.15
    r4 = p4.add_run("COMPUTER SCIENCE AND ENGINEERING")
    r4.font.name = "Comfortaa"
    r4.font.size = Pt(20)
    r4.font.bold = True
    r4._r.get_or_add_rPr().append(parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Comfortaa" w:hAnsi="Comfortaa" w:cs="Comfortaa"/>'))

    # Paragraph 5: ASSIGNMENT
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_before = Pt(0)
    p5.paragraph_format.space_after = Pt(4)
    p5.paragraph_format.line_spacing = 1.15
    r5 = p5.add_run("ASSIGNMENT")
    r5.font.name = "Comfortaa"
    r5.font.size = Pt(20)
    r5.font.bold = True
    r5._r.get_or_add_rPr().append(parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Comfortaa" w:hAnsi="Comfortaa" w:cs="Comfortaa"/>'))

    # Paragraph 6: Assignment Title
    p6 = doc.add_paragraph()
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p6.paragraph_format.space_before = Pt(0)
    p6.paragraph_format.space_after = Pt(46)
    p6.paragraph_format.line_spacing = 1.15
    r6 = p6.add_run(str(title).strip())
    r6.font.name = "Comfortaa"
    r6.font.size = Pt(20)
    r6.font.bold = True
    r6._r.get_or_add_rPr().append(parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Comfortaa" w:hAnsi="Comfortaa" w:cs="Comfortaa"/>'))

    # Paragraph 7: Course Code (14 pt as specified in requirement 5 & 6)
    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p7.paragraph_format.space_before = Pt(0)
    p7.paragraph_format.space_after = Pt(62)
    p7.paragraph_format.line_spacing = 1.15
    course_str = str(course).strip()
    course_display = course_str if course_str.upper().startswith("COURSE CODE:") else f"COURSE CODE: {course_str}"
    r7 = p7.add_run(course_display)
    r7.font.name = "Comfortaa"
    r7.font.size = Pt(14)
    r7.font.bold = True
    r7._r.get_or_add_rPr().append(parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Comfortaa" w:hAnsi="Comfortaa" w:cs="Comfortaa"/>'))

    # Lower Two-Column Section: Table
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(r"""
        <w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
    """)
    tblPr.append(tblBorders)

    row = table.rows[0]
    cell_left = row.cells[0]
    cell_right = row.cells[1]
    cell_left.width = Inches(3.35)
    cell_right.width = Inches(3.15)

    tcPr_left = cell_left._tc.get_or_add_tcPr()
    tcMar_left = parse_xml(r"""
        <w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:w="0" w:type="dxa"/>
            <w:left w:w="260" w:type="dxa"/>
            <w:bottom w:w="0" w:type="dxa"/>
            <w:right w:w="140" w:type="dxa"/>
        </w:tcMar>
    """)
    tcPr_left.append(tcMar_left)

    tcPr_right = cell_right._tc.get_or_add_tcPr()
    tcMar_right = parse_xml(r"""
        <w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:w="0" w:type="dxa"/>
            <w:left w:w="140" w:type="dxa"/>
            <w:bottom w:w="0" w:type="dxa"/>
            <w:right w:w="0" w:type="dxa"/>
        </w:tcMar>
    """)
    tcPr_right.append(tcMar_right)

    def add_cell_line(cell, text, size_pt=14, bold=False, color_rgb=(0,0,0), space_after=4):
        if len(cell.paragraphs) == 1 and cell.paragraphs[0].text == "":
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = "Comfortaa"
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*color_rgb)
        run._r.get_or_add_rPr().append(parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Comfortaa" w:hAnsi="Comfortaa" w:cs="Comfortaa"/>'))
        return p

    # Left Column: SUBMITTED BY (16pt heading, 14pt fields)
    add_cell_line(cell_left, "SUBMITTED BY", 16, bold=True, color_rgb=(67, 67, 67), space_after=10)
    add_cell_line(cell_left, f"Name: {sname}", 14, bold=False, color_rgb=(0, 0, 0), space_after=4)
    add_cell_line(cell_left, f"Student ID: {sid}", 14, bold=False, color_rgb=(0, 0, 0), space_after=4)
    add_cell_line(cell_left, f"Semester: {semester}", 14, bold=False, color_rgb=(0, 0, 0), space_after=0)

    # Right Column: SUBMITTED TO (16pt heading, 14pt fields)
    # Includes editable teacher designation and fixed college name below it
    add_cell_line(cell_right, "SUBMITTED TO", 16, bold=True, color_rgb=(67, 67, 67), space_after=10)
    add_cell_line(cell_right, f"Teacher Name: {tname}", 14, bold=False, color_rgb=(0, 0, 0), space_after=4)
    add_cell_line(cell_right, f"{tdesignation}", 14, bold=False, color_rgb=(0, 0, 0), space_after=4)
    add_cell_line(cell_right, "SHYAMOLI ENGINEERING COLLEGE", 14, bold=False, color_rgb=(0, 0, 0), space_after=0)

    doc.save(output_path)
    return output_path

def generate_pdf(title, course, sname, sid, semester, tname, tdesignation, output_path):
    c = canvas.Canvas(output_path, pagesize=letter)
    width, height = letter # 612 x 792 pt

    # Background Watermark
    if os.path.exists(WATERMARK_IMG):
        # Center x=306, y=374, size=400x400 (matches PDF exactly: [105.5, 174.4, 505.1, 574.0])
        c.drawImage(WATERMARK_IMG, 106, 174, width=399.5, height=399.5, mask='auto', preserveAspectRatio=True)

    # Heading: SHYAMOLI ENGINEERING COLLEGE
    font_title = 'Antonio-Bold' if 'Antonio-Bold' in pdfmetrics.getRegisteredFontNames() else 'Helvetica-Bold'
    font_sub = 'Times-Bold'
    font_body = 'Comfortaa-Regular' if 'Comfortaa-Regular' in pdfmetrics.getRegisteredFontNames() else 'Helvetica'
    font_bold = 'Comfortaa-Regular' if 'Comfortaa-Regular' in pdfmetrics.getRegisteredFontNames() else 'Helvetica-Bold'

    c.setFillColor(HexColor('#000000'))
    c.setFont(font_title, 43)
    c.drawCentredString(width / 2.0, 630.3, "SHYAMOLI ENGINEERING COLLEGE")

    # Institutional subtitle: A CONSTITUENT COLLEGE OF UNIVERSITY OF DHAKA
    c.setFillColor(HexColor('#077ACD'))
    c.setFont(font_sub, 14)
    c.drawCentredString(width / 2.0, 594.5, "A CONSTITUENT COLLEGE OF UNIVERSITY OF DHAKA")

    # DEPARTMENT OF
    c.setFillColor(HexColor('#000000'))
    c.setFont(font_bold, 20)
    c.drawCentredString(width / 2.0, 501.2, "DEPARTMENT OF")

    # COMPUTER SCIENCE AND ENGINEERING
    c.drawCentredString(width / 2.0, 473.9, "COMPUTER SCIENCE AND ENGINEERING")

    # ASSIGNMENT
    c.drawCentredString(width / 2.0, 401.6, "ASSIGNMENT")

    # Assignment Title
    c.drawCentredString(width / 2.0, 374.3, str(title).strip())

    # Course Code
    course_str = str(course).strip()
    course_display = course_str if course_str.upper().startswith("COURSE CODE:") else f"COURSE CODE: {course_str}"
    c.setFont(font_bold, 14)
    c.drawCentredString(width / 2.0, 305.5, course_display)

    # SUBMITTED BY (x = 90.3 pt)
    c.setFillColor(HexColor('#434343'))
    c.setFont(font_bold, 16)
    c.drawString(90.3, 216.3, "SUBMITTED BY")

    c.setFillColor(HexColor('#000000'))
    c.setFont(font_body, 14)
    c.drawString(90.3, 190.0, f"Name: {sname}")
    c.drawString(90.3, 170.0, f"Student ID: {sid}")
    c.drawString(90.3, 150.0, f"Semester: {semester}")

    # SUBMITTED TO (x = 345.3 pt)
    c.setFillColor(HexColor('#434343'))
    c.setFont(font_bold, 16)
    c.drawString(345.3, 216.3, "SUBMITTED TO")

    c.setFillColor(HexColor('#000000'))
    c.setFont(font_body, 14)
    c.drawString(345.3, 190.0, f"Teacher Name: {tname}")
    c.drawString(345.3, 170.0, str(tdesignation).strip())
    c.drawString(345.3, 150.0, "SHYAMOLI ENGINEERING COLLEGE")

    c.showPage()
    c.save()
    return output_path
